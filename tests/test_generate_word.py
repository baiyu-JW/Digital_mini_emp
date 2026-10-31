"""文档生成工具单元测试（pytest）。

运行：在项目根目录执行  python -m pytest tests/ -q
前置：pip install python-docx
"""
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from tools.generate_word import generate_word, WordGenError
from docx import Document


def test_generate_word_basic(tmp_path):
    out = tmp_path / "周报.docx"
    p = generate_word(
        "第一段正文。\n第二段正文。\n# 小结\n本周表现平稳。",
        str(out),
        title="测试周报",
        meta={"统计周期": "2026年第44周"},
    )
    assert Path(p).exists()
    doc = Document(p)
    texts = [para.text for para in doc.paragraphs]
    assert "测试周报" in texts
    assert "统计周期：2026年第44周" in texts
    assert "第一段正文。" in texts
    assert any("小结" in t for t in texts)


def test_generate_word_empty_raises(tmp_path):
    import pytest
    with pytest.raises(WordGenError):
        generate_word("", str(tmp_path / "x.docx"))


def test_generate_word_idempotent(tmp_path):
    """重复调用覆盖同一文件，不产生多个副本（工程加固：幂等）。"""
    out = tmp_path / "周报.docx"
    generate_word("第一版", str(out))
    generate_word("第二版", str(out))
    doc = Document(out)
    joined = "\n".join(para.text for para in doc.paragraphs)
    assert "第一版" not in joined
    assert "第二版" in joined
