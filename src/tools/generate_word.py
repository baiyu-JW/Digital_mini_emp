"""文档生成工具：把周报文本生成带格式的 Word（赛题场景 A 交付物）。

Day 4 初版：先把字写进去（标题/元信息/正文/简单 Markdown 标题）。
图表接口已预留（charts 参数），Day 6 接入 matplotlib 后传入即可。
设计原则（对应赛题"工程加固"加分）：
- 入参/出参明确，失败抛自定义异常而非静默。
- 写文件基于 out_path 确定性命名，重复调用覆盖同一文件（幂等，不产生一堆副本）。
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt


class WordGenError(Exception):
    """Word 生成失败的统一异常。"""


def generate_word(report_text, out_path, title="经营周报", charts=None, meta=None):
    """生成带格式的 Word 周报。

    入参:
        report_text (str): 周报正文。按换行拆段落；以 '#' 开头的行当作 Markdown 标题。
        out_path (str):     输出 .docx 路径（父目录自动创建）。
        title (str):        文档主标题。
        charts (list[str]|None): 图表 png 绝对路径列表（Day 6 接入；本版存在则追加）。
        meta (dict|None):   元信息，如 {"统计周期": "2026年第44周"}，写入首段斜体。
    出参:
        str: 实际写出的文件路径。
    异常:
        WordGenError: 任何写出/渲染失败。
    """
    if not report_text or not report_text.strip():
        raise WordGenError("report_text 为空，无法生成周报")

    try:
        doc = Document()

        doc.add_heading(title, level=0)

        if meta:
            for k, v in meta.items():
                p = doc.add_paragraph()
                run = p.add_run(f"{k}：{v}")
                run.italic = True

        for block in report_text.split("\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("#"):
                level = min(block.count("#", 0, 6), 3) or 1
                doc.add_heading(block.lstrip("#").strip(), level=level)
            else:
                doc.add_paragraph(block)

        if charts:
            for c in charts:
                if c and Path(c).exists():
                    doc.add_picture(c)

        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        return str(out)
    except WordGenError:
        raise
    except Exception as e:  # noqa: BLE001 - 统一包装，避免 docx 内部异常外泄
        raise WordGenError(f"Word 生成失败：{e}") from e
